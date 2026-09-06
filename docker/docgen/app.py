"""docgen - document toolbox for Automation Lab workflows.

Endpoints
---------
GET  /health                       -> versions + installed OCR languages
POST /render/pdf   (JSON)          -> HTML -> PDF (WeasyPrint). Body: {html, css?, filename?}
POST /render/docx  (JSON)          -> structured document -> DOCX (Arabic RTL aware)
POST /render/pptx  (JSON)          -> structured document -> PPTX (Arabic RTL aware)
POST /pdf/extract  (multipart)     -> text + tables per page (pdfplumber)
POST /ocr?lang=eng|ara|ara+eng     -> text, lines and confidence (Tesseract)

All responses for binary formats carry Content-Disposition with the requested filename so the n8n
HTTP Request node (Response Format: File) keeps a sensible file name.
"""
from __future__ import annotations

import io
import re
import subprocess
from typing import Any

from fastapi import FastAPI, File, HTTPException, Query, UploadFile
from fastapi.responses import JSONResponse, Response
from pydantic import BaseModel, Field

app = FastAPI(title="Automation Lab docgen", version="1.0.0")

FONT_DIR = "/app/fonts"
BASE_CSS = f"""
@font-face {{ font-family: 'Amiri'; src: url('file://{FONT_DIR}/Amiri-Regular.ttf'); }}
@font-face {{ font-family: 'Amiri'; font-weight: bold; src: url('file://{FONT_DIR}/Amiri-Bold.ttf'); }}
@page {{ size: A4; margin: 18mm 16mm; }}
body {{ font-family: 'DejaVu Sans', 'Amiri', sans-serif; font-size: 11pt; color: #1a1a1a; }}
[dir="rtl"], .rtl {{ direction: rtl; unicode-bidi: embed; font-family: 'Amiri', 'DejaVu Sans', serif; }}
table {{ border-collapse: collapse; width: 100%; }}
th, td {{ border: 1px solid #cfcfcf; padding: 4px 6px; text-align: start; }}
th {{ background: #f2f2f2; }}
"""

ARABIC_RE = re.compile(r"[؀-ۿ]")


def _filename_header(filename: str) -> dict[str, str]:
    safe = re.sub(r"[^A-Za-z0-9._-]", "_", filename) or "document"
    return {"Content-Disposition": f'attachment; filename="{safe}"'}


# ----------------------------------------------------------------------------- health ----
@app.get("/health")
def health() -> dict[str, Any]:
    langs: list[str] = []
    try:
        out = subprocess.run(["tesseract", "--list-langs"], capture_output=True, text=True, timeout=10).stdout
        langs = [l.strip() for l in out.splitlines()[1:] if l.strip()]
    except Exception:  # noqa: BLE001
        pass
    try:
        import weasyprint  # type: ignore

        wp = weasyprint.__version__
    except Exception:  # noqa: BLE001
        wp = "unavailable"
    return {"status": "ok", "service": "docgen", "weasyprint": wp, "tesseract_langs": langs,
            "fonts": ["Amiri", "DejaVu Sans"]}


# --------------------------------------------------------------------------- HTML -> PDF ----
class PdfRequest(BaseModel):
    html: str = Field(..., description="Full or partial HTML document")
    css: str | None = Field(None, description="Extra CSS appended after the base stylesheet")
    filename: str = "document.pdf"


@app.post("/render/pdf")
def render_pdf(req: PdfRequest) -> Response:
    try:
        from weasyprint import CSS, HTML  # type: ignore
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(500, f"weasyprint unavailable: {exc}") from exc
    html = req.html
    if "<html" not in html.lower():
        rtl = ' dir="rtl" lang="ar"' if ARABIC_RE.search(html) else ""
        html = f"<!doctype html><html{rtl}><head><meta charset='utf-8'></head><body>{html}</body></html>"
    stylesheets = [CSS(string=BASE_CSS + (req.css or ""))]
    pdf = HTML(string=html, base_url="/app").write_pdf(stylesheets=stylesheets)
    return Response(content=pdf, media_type="application/pdf", headers=_filename_header(req.filename))


# ----------------------------------------------------------------- structured documents ----
class Table(BaseModel):
    columns: list[str]
    rows: list[list[Any]] = []


class Section(BaseModel):
    heading: str | None = None
    paragraphs: list[str] = []
    bullets: list[str] = []
    table: Table | None = None


class DocRequest(BaseModel):
    title: str
    subtitle: str | None = None
    lang: str = Field("en", description="'ar' or 'en' (controls RTL and default font)")
    rtl: bool | None = Field(None, description="Force text direction; defaults to lang == 'ar'")
    sections: list[Section] = []
    footer: str | None = None
    filename: str | None = None


def _is_rtl(req: DocRequest) -> bool:
    return req.rtl if req.rtl is not None else req.lang.lower().startswith("ar")


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{v:,.2f}"
    return str(v)


@app.post("/render/docx")
def render_docx(req: DocRequest) -> Response:
    try:
        from docx import Document  # type: ignore
        from docx.enum.table import WD_TABLE_ALIGNMENT  # type: ignore
        from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
        from docx.oxml import OxmlElement  # type: ignore
        from docx.oxml.ns import qn  # type: ignore
        from docx.shared import Pt  # type: ignore
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(500, f"python-docx unavailable: {exc}") from exc

    rtl = _is_rtl(req)
    font_name = "Amiri" if rtl else "Calibri"
    doc = Document()

    # Default font for Latin and complex scripts (Arabic uses the 'cs' slot).
    style = doc.styles["Normal"]
    style.font.name = font_name
    style.font.size = Pt(12 if rtl else 11)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), font_name)

    def mark_rtl(paragraph) -> None:
        if not rtl:
            return
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        ppr = paragraph._p.get_or_add_pPr()
        bidi = OxmlElement("w:bidi")
        bidi.set(qn("w:val"), "1")
        ppr.append(bidi)
        for run in paragraph.runs:
            rp = run._r.get_or_add_rPr()
            r = OxmlElement("w:rtl")
            r.set(qn("w:val"), "1")
            rp.append(r)
            run.font.name = font_name
            rf = rp.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts")
                rp.append(rf)
            rf.set(qn("w:cs"), font_name)

    h = doc.add_heading(req.title, level=0)
    mark_rtl(h)
    if req.subtitle:
        p = doc.add_paragraph(req.subtitle)
        mark_rtl(p)

    for section in req.sections:
        if section.heading:
            mark_rtl(doc.add_heading(section.heading, level=1))
        for text in section.paragraphs:
            mark_rtl(doc.add_paragraph(text))
        for text in section.bullets:
            mark_rtl(doc.add_paragraph(text, style="List Bullet"))
        if section.table:
            cols = section.table.columns
            table = doc.add_table(rows=1, cols=len(cols))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            if rtl:
                tblpr = table._tbl.tblPr
                bidi = OxmlElement("w:bidiVisual")
                bidi.set(qn("w:val"), "1")
                tblpr.append(bidi)
            for i, col in enumerate(cols):
                cell = table.rows[0].cells[i]
                cell.text = str(col)
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.bold = True
                    mark_rtl(para)
            for row in section.table.rows:
                cells = table.add_row().cells
                for i in range(len(cols)):
                    cells[i].text = _cell(row[i] if i < len(row) else "")
                    for para in cells[i].paragraphs:
                        mark_rtl(para)
            doc.add_paragraph("")

    if req.footer:
        mark_rtl(doc.add_paragraph(req.footer))

    buf = io.BytesIO()
    doc.save(buf)
    name = req.filename or ("report-ar.docx" if rtl else "report.docx")
    return Response(content=buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    headers=_filename_header(name))


@app.post("/render/pptx")
def render_pptx(req: DocRequest) -> Response:
    try:
        from pptx import Presentation  # type: ignore
        from pptx.enum.text import PP_ALIGN  # type: ignore
        from pptx.util import Inches, Pt  # type: ignore
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(500, f"python-pptx unavailable: {exc}") from exc

    rtl = _is_rtl(req)
    font_name = "Amiri" if rtl else "Calibri"
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def style_frame(tf, size: int = 18) -> None:
        for para in tf.paragraphs:
            if rtl:
                para.alignment = PP_ALIGN.RIGHT
                para._p.get_or_add_pPr().set("rtl", "1")
            for run in para.runs:
                run.font.name = font_name
                run.font.size = Pt(size)

    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = req.title
    style_frame(title_slide.shapes.title.text_frame, 40)
    if req.subtitle and len(title_slide.placeholders) > 1:
        title_slide.placeholders[1].text = req.subtitle
        style_frame(title_slide.placeholders[1].text_frame, 22)

    blank = prs.slide_layouts[5]  # title only
    for section in req.sections:
        slide = prs.slides.add_slide(blank)
        slide.shapes.title.text = section.heading or req.title
        style_frame(slide.shapes.title.text_frame, 30)
        top = Inches(1.5)
        body_lines = list(section.paragraphs) + [f"• {b}" for b in section.bullets]
        if body_lines:
            box = slide.shapes.add_textbox(Inches(0.6), top, Inches(12.1), Inches(2.2))
            tf = box.text_frame
            tf.word_wrap = True
            tf.text = body_lines[0]
            for line in body_lines[1:]:
                tf.add_paragraph().text = line
            style_frame(tf, 18)
            top = Inches(3.9)
        if section.table:
            cols = section.table.columns
            rows = section.table.rows
            shape = slide.shapes.add_table(len(rows) + 1, len(cols), Inches(0.6), top, Inches(12.1),
                                           Inches(0.4) * (len(rows) + 1))
            table = shape.table
            order = list(range(len(cols)))
            if rtl:
                order.reverse()  # visual right-to-left column order
            for vi, ci in enumerate(order):
                table.cell(0, vi).text = str(cols[ci])
                style_frame(table.cell(0, vi).text_frame, 14)
            for r, row in enumerate(rows, start=1):
                for vi, ci in enumerate(order):
                    table.cell(r, vi).text = _cell(row[ci] if ci < len(row) else "")
                    style_frame(table.cell(r, vi).text_frame, 13)

    if req.footer:
        slide = prs.slides.add_slide(blank)
        slide.shapes.title.text = req.footer
        style_frame(slide.shapes.title.text_frame, 24)

    buf = io.BytesIO()
    prs.save(buf)
    name = req.filename or ("report-ar.pptx" if rtl else "report.pptx")
    return Response(content=buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    headers=_filename_header(name))


# ---------------------------------------------------------------------- PDF extraction ----
@app.post("/pdf/extract")
async def pdf_extract(file: UploadFile = File(...), tables: bool = Query(True)) -> JSONResponse:
    try:
        import pdfplumber  # type: ignore
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(500, f"pdfplumber unavailable: {exc}") from exc
    raw = await file.read()
    if not raw:
        raise HTTPException(400, "empty file")
    pages: list[dict[str, Any]] = []
    flat_tables: list[dict[str, Any]] = []
    try:
        with pdfplumber.open(io.BytesIO(raw)) as pdf:
            for idx, page in enumerate(pdf.pages, start=1):
                text = page.extract_text() or ""
                page_tables: list[list[list[str]]] = []
                if tables:
                    for t in page.extract_tables() or []:
                        cleaned = [[(c or "").strip() for c in row] for row in t if any((c or "").strip() for c in row)]
                        if cleaned:
                            page_tables.append(cleaned)
                            flat_tables.append({"page": idx, "header": cleaned[0], "rows": cleaned[1:]})
                pages.append({"page": idx, "text": text, "tables": page_tables})
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(422, f"could not parse PDF: {exc}") from exc
    return JSONResponse({"filename": file.filename, "pages": len(pages), "text": "\n\n".join(p["text"] for p in pages),
                         "tables": flat_tables, "page_data": pages})


# ------------------------------------------------------------------------------- OCR ----
@app.post("/ocr")
async def ocr(file: UploadFile = File(...), lang: str = Query("eng"), psm: int = Query(6)) -> JSONResponse:
    try:
        import pytesseract  # type: ignore
        from PIL import Image, ImageOps  # type: ignore
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(500, f"OCR stack unavailable: {exc}") from exc
    raw = await file.read()
    if not raw:
        raise HTTPException(400, "empty file")
    try:
        img = Image.open(io.BytesIO(raw))
        img = ImageOps.exif_transpose(img).convert("L")
        if img.width < 1200:
            scale = 1200 / img.width
            img = img.resize((int(img.width * scale), int(img.height * scale)))
        img = ImageOps.autocontrast(img)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(422, f"not an image: {exc}") from exc
    lang = re.sub(r"[^a-z+]", "", lang.lower()) or "eng"
    config = f"--psm {int(psm)}"
    try:
        data = pytesseract.image_to_data(img, lang=lang, config=config, output_type=pytesseract.Output.DICT)
        text = pytesseract.image_to_string(img, lang=lang, config=config)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(500, f"tesseract failed: {exc}") from exc
    lines: dict[tuple[int, int, int], list[str]] = {}
    confs: list[float] = []
    for i, word in enumerate(data["text"]):
        if not str(word).strip():
            continue
        key = (data["block_num"][i], data["par_num"][i], data["line_num"][i])
        lines.setdefault(key, []).append(str(word))
        try:
            c = float(data["conf"][i])
            if c >= 0:
                confs.append(c)
        except (TypeError, ValueError):
            pass
    line_list = [" ".join(w) for _, w in sorted(lines.items())]
    return JSONResponse({"filename": file.filename, "lang": lang, "text": text.strip(), "lines": line_list,
                         "words": len(confs), "mean_confidence": round(sum(confs) / len(confs), 1) if confs else 0.0})
